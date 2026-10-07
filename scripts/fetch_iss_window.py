"""Read-only: download IRAS Issue (ISS) transactions for one time range and print them.

    python scripts/fetch_iss_window.py 2026-10-06 17:30 18:30

Reuses the scraper's own login and export code. Writes only to <data>/ISS_adhoc/ and
touches no database. Run under the daily_scrape lock (see the shell line in the docs):
    flock -w 900 /data/locks/daily_scrape.lock python scripts/fetch_iss_window.py ...
"""
import asyncio
import sys
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scrapers"))

import openpyxl  # noqa: E402
from playwright.async_api import async_playwright  # noqa: E402

import daily_scrape as ds  # noqa: E402

_iss = ds._iss


def windows(day, start, end):
    t = datetime.strptime(f"{day} {start}", "%Y-%m-%d %H:%M")
    stop = datetime.strptime(f"{day} {end}", "%Y-%m-%d %H:%M")
    while t < stop:
        n = t + timedelta(minutes=30)
        yield (t.strftime("%Y-%m-%d"), t.hour, t.minute, n.strftime("%Y-%m-%d"), n.hour, n.minute)
        t = n


def summarise(files, big=100.0):
    """Per window: rows and litres by product; then every sale of `big` litres or more."""
    bigs = []
    for f in sorted(files):
        ws = openpyxl.load_workbook(f, read_only=True, data_only=True).active
        rows = ws.iter_rows(values_only=True)
        head = list(next(rows))
        ix = {h: i for i, h in enumerate(head)}
        litres, n, mx = {}, 0, 0.0
        for r in rows:
            if r[0] is None:
                continue
            n += 1
            q = float(r[ix["Quantity"]] or 0)
            litres[r[ix["Product Code"]]] = litres.get(r[ix["Product Code"]], 0.0) + q
            if q >= big:
                bigs.append((r[ix["Transaction Start Time"]], r[ix["Transaction Date"]], r[ix["Product Code"]], q,
                             r[ix["Amount"]], r[ix["Nozzle No"]], r[ix["Vehicle Number"]], r[ix["Attendant ID"]],
                             r[ix["Pmt Transaction Code"]], r[ix["ISS Transaction Type"]]))
        print(f"{f.name[4:-5]}: {n:3d} sales  " + "  ".join(f"{k} {v:,.0f}L" for k, v in sorted(litres.items())))
    print(f"\nSales of {big:g} L or more:")
    for b in sorted(bigs, key=lambda x: str(x[0])):
        print("  start", b[0], b[1], b[2], f"{b[3]:,.1f}L", f"Rs{b[4]}", "nozzle", b[5], "veh", b[6], "att", b[7], "pmt", b[8], b[9])


async def main(day, start, end):
    out_dir = ds._data_root / "ISS_adhoc"
    out_dir.mkdir(parents=True, exist_ok=True)
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True, args=["--no-sandbox", "--disable-dev-shm-usage"])
        kw = {
            "accept_downloads": True,
            "viewport": {"width": 1400, "height": 900},
            # same browser identity as daily_scrape: the portal renders blank for "HeadlessChrome"
            "user_agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                           "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"),
        }
        proxy = ds.iras_proxy_cfg()
        if proxy is not None:
            kw["proxy"] = proxy
        ctx = await browser.new_context(**kw)
        page = await ctx.new_page()
        await page.goto(ds.LOGIN_URL, wait_until="networkidle", timeout=30_000)
        await page.wait_for_timeout(1000)
        if not await ds._autonomous_login(page, debug_dir=ds._make_login_debug_dir()):
            sys.exit("login failed")
        await _iss.navigate_to_iss(page, shift_date=day)
        await _iss.ensure_iss_archive_mode(page, day)
        files = []
        for w in windows(day, start, end):
            await _iss.handle_session_expiry(page)
            res = await _iss.export_window(page, out_dir, *w)
            print(f"window {w[0]} {w[1]:02d}:{w[2]:02d}-{w[4]:02d}:{w[5]:02d} -> {res}")
            if res is True:
                files.append(out_dir / _iss.filename_safe(*w))
            await asyncio.sleep(_iss.DELAY_BETWEEN)
        await browser.close()
    if len(sys.argv) > 4 and sys.argv[4] == "summary":
        summarise(files)
        return
    for f in files:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        print(f"\n=== {f.name} ===")
        for row in wb.active.iter_rows(values_only=True):
            print(" | ".join("" if c is None else str(c) for c in row))


if __name__ == "__main__":
    asyncio.run(main(*sys.argv[1:4]))
